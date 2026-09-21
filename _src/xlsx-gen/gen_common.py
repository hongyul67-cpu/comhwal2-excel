# -*- coding: utf-8 -*-
"""컴활 학습용 엑셀 예제파일 - 공통 유틸"""
import sys, re, pathlib
sys.stdout.reconfigure(encoding='utf-8')

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "1F3864"
BLUE = "2E5C8A"
YEL  = "FFF2CC"
GREY = "F2F2F2"
GREEN= "E2EFDA"
PINK = "FCE4EC"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def astext(c):
    """설명 글이 = 로 시작하면 수식이 아니라 '글자' 로 저장한다"""
    if isinstance(c.value, str) and c.value.startswith("="):
        c.data_type = "s"
    return c


def hdr(ws, row, values, start_col=1, fill=BLUE):
    for i, v in enumerate(values):
        c = astext(ws.cell(row=row, column=start_col + i, value=v))
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 24

def put(ws, row, values, start_col=1, fill=None, fmt=None, align=None, bold=False):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=v)
        if isinstance(v, str) and v.startswith("="):
            c.data_type = "s"          # 수식이 아니라 "수식을 적은 글자"
        c.border = BORDER
        if fill:  c.fill = PatternFill("solid", fgColor=fill)
        if fmt:   c.number_format = fmt
        if bold:  c.font = Font(bold=True)
        if align: c.alignment = Alignment(horizontal=align, vertical="center")
        else:     c.alignment = Alignment(vertical="center")

def blank(ws, row, cols, fill=YEL):
    """학생이 채울 노란 칸"""
    for col in cols:
        c = ws.cell(row=row, column=col)
        c.fill = PatternFill("solid", fgColor=fill)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")

def title(ws, row, text, span=8, size=14, color=NAVY):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = astext(ws.cell(row=row, column=1, value=text))
    c.font = Font(bold=True, size=size, color=color)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 26

def note(ws, row, text, span=8, color="806000", fill="FFF8E1"):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = astext(ws.cell(row=row, column=1, value=text))
    c.font = Font(size=10, color=color)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c.fill = PatternFill("solid", fgColor=fill)
    ws.row_dimensions[row].height = 30

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def guide_sheet(wb, lines, sheet_title="읽어보기"):
    ws = wb.create_sheet(sheet_title, 0)
    widths(ws, {"A": 4, "B": 96})
    ws.sheet_view.showGridLines = False
    r = 2
    for kind, text in lines:
        c = astext(ws.cell(row=r, column=2, value=text))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if kind == "h1":
            c.font = Font(bold=True, size=16, color=NAVY); ws.row_dimensions[r].height = 30
        elif kind == "h2":
            c.font = Font(bold=True, size=12, color=BLUE); ws.row_dimensions[r].height = 24
        elif kind == "note":
            c.font = Font(size=10, color="806000")
            c.fill = PatternFill("solid", fgColor="FFF8E1")
            ws.row_dimensions[r].height = 34
        else:
            c.font = Font(size=11); ws.row_dimensions[r].height = 19
        r += 1
    return ws

def mission_sheet(wb, name, head, rows, note_text=None, colw=None):
    """번호 | 미션 | (빈칸) | 힌트 형태의 미션표"""
    ws = wb.create_sheet(name)
    widths(ws, colw or {"A": 5, "B": 62, "C": 16, "D": 34})
    title(ws, 1, head, span=4)
    if note_text:
        note(ws, 2, note_text, span=4)
    hdr(ws, 4, ["번호", "미션 — 무엇을 하는가", "답(여기에 입력)", "힌트 (쓰는 함수·메뉴)"])
    r = 5
    for i, (m, hint) in enumerate(rows, 1):
        put(ws, r, [i, m, None, hint])
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=2).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row=r, column=4).alignment = Alignment(vertical="center", wrap_text=True)
        blank(ws, r, [3])
        ws.row_dimensions[r].height = 30
        r += 1
    return ws

def answer_sheet(wb, name, head, rows, colw=None):
    ws = wb.create_sheet(name)
    widths(ws, colw or {"A": 5, "B": 46, "C": 48, "D": 40})
    title(ws, 1, head, span=4)
    note(ws, 2, "혼자 공부할 때만 펼쳐 보세요. 먼저 스스로 만들어 보고, 막혔을 때 확인하는 칸입니다.", span=4)
    hdr(ws, 4, ["번호", "미션", "정답 수식 / 순서", "왜 이렇게 되는가"])
    r = 5
    for i, (m, a, why) in enumerate(rows, 1):
        put(ws, r, [i, m, a, why], fill=GREEN if i % 2 == 0 else None)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
        for col in (2, 3, 4):
            ws.cell(row=r, column=col).alignment = Alignment(vertical="center", wrap_text=True)
        ac = ws.cell(row=r, column=3)
        ac.data_type = "s"
        ac.font = Font(name="Consolas", size=10, bold=True)
        ws.row_dimensions[r].height = 32
        r += 1
    return ws

def apply_print(wb, landscape=True):
    for ws in wb.worksheets:
        ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.3
        ws.page_margins.top = ws.page_margins.bottom = 0.4

def save(wb, path):
    wb.calculation.fullCalcOnLoad = True
    apply_print(wb)
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    p = pathlib.Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    print("saved:", p, "| sheets:", wb.sheetnames)
    return p


def lit(ws, row, col, s, mono=True, bold=False, color=None):
    """수식처럼 생긴 글자를 '글자 그대로' 넣기"""
    c = ws.cell(row=row, column=col, value=s)
    c.data_type = "s"
    c.border = BORDER
    c.font = Font(name="Consolas" if mono else "맑은 고딕", size=10, bold=bold, color=color)
    c.alignment = Alignment(vertical="center")
    return c
