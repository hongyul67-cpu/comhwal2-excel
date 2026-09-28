# -*- coding: utf-8 -*-
"""함수 연습소(웹)와 똑같은 문제를 수업 범위별 엑셀 파일로 굽는다.

  node export_web.js      → out/web_problems.json   (웹 데이터 그대로)
  python gen_fx.py out    → out/fx/2급_함수연습_01_논리_v1.xlsx … 8개
  node export_web.js 1 · python gen_fx.py out 1   → 1급 out/fx1/1급_함수연습_… 10개

· 문제 하나 = 시트 하나. 표는 A1 부터 놓아 **웹 화면과 칸 주소가 똑같다** (웹에서 E2 면 엑셀도 E2).
· 여러 명 표는 첫 칸에 수식을 넣고 채우기 핸들로 끌어내린다. 칸마다 채점한다.
· 숨은 AA열 = 그 칸의 모범답안(끌어내린 모양 그대로), AB열 = 그 칸이 맞았는지.
· 학생이 한 명씩 내려받아 각자 쓰는 파일이다.
"""
import sys, os, json, re, datetime as dt
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gen_lib import *
from openpyxl.worksheet.formula import ArrayFormula

CATS = [  # 웹 범위 이름, 파일 키, 짧은 이름
    ("논리(IF)", "logic", "논리"),
    ("통계", "stat", "통계"),
    ("조건부 집계", "cond", "조건부집계"),
    ("수학·반올림", "math", "수학반올림"),
    ("문자열", "text", "문자열"),
    ("찾기·참조", "lookup", "찾기참조"),
    ("데이터베이스", "db", "DB함수"),
    ("날짜", "date", "날짜"),
]
CATS1 = [
    ("논리(IF)", "logic", "논리"), ("통계", "stat", "통계"), ("조건부 집계", "cond", "조건부집계"),
    ("찾기·참조", "lookup", "찾기참조"), ("문자열", "text", "문자열"), ("수학·반올림", "math", "수학반올림"),
    ("배열수식", "array", "배열수식"), ("날짜", "date", "날짜"), ("다중 조건(IFS)", "ifs", "다중조건"),
    ("데이터베이스", "db", "DB함수"),
]
GRADE = "2"
VER = "v1"
ANS_COL, CHK_COL = 27, 28          # AA · AB (숨김)
INFO = 8                           # H열부터 문제 설명
DATE_RE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")
TIME_RE = re.compile(r"^(\d{2}):(\d{2}):(\d{2})$")


def dlname(i, short):
    return "%s급_함수연습_%02d_%s_%s.xlsx" % (GRADE, i, short, VER)


def cellval(v):
    """웹의 '2024-07-16' · '14:35:00' 은 엑셀에서는 진짜 날짜·시각으로"""
    if isinstance(v, str):
        m = DATE_RE.match(v)
        if m:
            return dt.date(int(m[1]), int(m[2]), int(m[3])), "yyyy-mm-dd"
        m = TIME_RE.match(v)
        if m:
            return dt.time(int(m[1]), int(m[2]), int(m[3])), "hh:mm:ss"
    return v, None


def sheet_name(no, title):
    t = re.sub(r'[\[\]:*?/\\\'"()]', "", title).strip()
    return ("%02d %s" % (no, t))[:28]


def row_check(addr, ans):
    """그 칸이 수식이고 · 값이 모범답안과 같은가 (글자는 대소문자까지 EXACT)"""
    return ("=AND(ISFORMULA(%s),IFERROR(ROUND(%s,4)=ROUND(%s,4),EXACT(%s,%s)))"
            % (addr, addr, ans, addr, ans))


def problem_sheet(wb, no, p):
    ws = wb.create_sheet(sheet_name(no, p["title"]))
    ws.sheet_view.showGridLines = True
    grid = p["grid"]
    ncol = max(len(r) for r in grid)
    tcol = p["c"]
    trows = [x["r"] for x in p["rows"]]
    first, last = "%s%d" % (p["col"], trows[0]), "%s%d" % (p["col"], trows[-1])
    multi = len(trows) > 1

    # ── 표 (A1 부터 — 웹과 같은 주소)
    for r, row in enumerate(grid, 1):
        for c in range(1, ncol + 1):
            v = row[c - 1] if c - 1 < len(row) else None
            cell = ws.cell(row=r, column=c)
            if c == tcol and r in trows:
                ph(ws, r, c)
                continue
            if v is None or v == "":
                if r == 1:
                    cell.border = BORDER
                continue
            val, nf = cellval(v)
            cell.value = val
            astext(cell)
            if nf:
                cell.number_format = nf
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if r == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=BLUE)
    for c in range(1, ncol + 1):
        head = str(grid[0][c - 1]) if c - 1 < len(grid[0]) and grid[0][c - 1] else ""
        ws.column_dimensions[get_column_letter(c)].width = max(11, min(22, len(head) * 2 + 6))
    ws.column_dimensions[get_column_letter(tcol)].width = 14
    ws.column_dimensions[get_column_letter(ncol + 1)].width = 3

    # ── 숨은 채점 기준 (배열 수식 문제는 엑셀에서도 배열 수식으로 — Ctrl+Shift+Enter 로 넣은 것과 같게)
    arr = p["cat"] == "배열수식"
    for x in p["rows"]:
        a = "%s%d" % (get_column_letter(ANS_COL), x["r"])
        ws[a] = ArrayFormula(a, fx(x["formula"])) if arr else fx(x["formula"])
        ws.cell(row=x["r"], column=CHK_COL,
                value=fx(row_check("%s%d" % (p["col"], x["r"]), "%s%d" % (get_column_letter(ANS_COL), x["r"]))))
    hide_helpers(ws, ANS_COL, CHK_COL)

    # ── 오른쪽 설명 (H열부터 — 표를 가리지 않게)
    H = max(INFO, ncol + 2)
    span = H + 6
    for c in range(H, span + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12

    def box(r, text, *, h=None, font=None, fill=None, rows=1):
        ws.merge_cells(start_row=r, start_column=H, end_row=r + rows - 1, end_column=span)
        c = astext(ws.cell(row=r, column=H, value=text))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if font: c.font = font
        if fill: c.fill = PatternFill("solid", fgColor=fill)
        if h:
            for k in range(rows):
                ws.row_dimensions[r + k].height = max(ws.row_dimensions[r + k].height or 15, h)
        return c

    box(1, "%d번 · %s" % (no, p["title"]), font=Font(bold=True, size=13, color=NAVY), h=24)
    box(2, p["prompt"], font=Font(size=11), fill="EEF3FA", rows=3, h=22)
    how = ("① %s 의 0 을 지우고 수식을 넣습니다.\n② 채우기 핸들(칸 오른쪽 아래 ■)을 %s 까지 끌어내립니다."
           % (first, last)) if multi else ("%s 의 0 을 지우고 수식을 넣습니다." % first)
    box(5, how, font=Font(size=10.5, color=ORANGE, bold=True), rows=2 if multi else 1, h=20)
    r = 7 if multi else 6
    lab = ws.cell(row=r, column=H, value="채점")
    lab.font = Font(bold=True, color="FFFFFF"); lab.fill = PatternFill("solid", fgColor=BLUE)
    lab.alignment = Alignment(horizontal="center", vertical="center")
    rng = "%s%d:%s%d" % (get_column_letter(CHK_COL), trows[0], get_column_letter(CHK_COL), trows[-1])
    n = len(trows)
    chk = ('=IF(NOT(ISFORMULA(%s)),"↖ 0 지우고 수식",IF(COUNTIF(%s,TRUE)=%d,"✔ 맞음",'
           '"✘ 다시 ("&COUNTIF(%s,TRUE)&" / %d칸)"))' % (first, rng, n, rng, n)) if multi else \
          ('=IF(NOT(ISFORMULA(%s)),"↖ 0 지우고 수식",IF(COUNTIF(%s,TRUE)=1,"✔ 맞음","✘ 다시"))' % (first, rng))
    ws.merge_cells(start_row=r, start_column=H + 1, end_row=r, end_column=H + 3)
    cc = ws.cell(row=r, column=H + 1, value=fx(chk))
    cc.font = Font(bold=True, size=12)
    cc.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(H, H + 4):
        ws.cell(row=r, column=c).border = BORDER
    mark_rules(ws, "%s%d" % (get_column_letter(H + 1), r))
    ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 24)
    if multi:
        box(r + 1, "칸마다 따로 채점합니다. 첫 칸만 맞고 아래가 틀리면 — 끌어내릴 때 움직이면 안 되는 주소에 $ 가 있는지 보세요.",
            font=Font(size=9.5, color="7F7F7F"), rows=2, h=18)
        r += 2
    if p["hint"]:
        box(r + 2, "💡 막히면: " + p["hint"], font=Font(size=10, color="7F7F7F"), rows=2, h=18)
    ws.freeze_panes = None
    return ws, "%s%d" % (get_column_letter(H + 1), 7 if multi else 6), first, last


def build(cat, idx, key, short, probs, outdir):
    wb = Workbook()
    guide_sheet(wb, [
        ("h1", "컴활 %s급 실기 · 함수 연습 — %s (%d문제)" % (GRADE, short, len(probs))),
        ("p", "웹 「함수 연습소」의 [%s] 범위와 똑같은 문제입니다. 번호도, 표도, 칸 주소도 같습니다." % cat),
        ("p", "웹에서는 한 문제씩, 여기서는 진짜 엑셀로 한꺼번에 풀어 봅니다."),
        ("h2", "푸는 법"),
        ("p", "① 아래 시트 탭(01, 02 …)이 문제 하나씩입니다. 오른쪽에 문제가 적혀 있습니다."),
        ("p", "② 노란 칸의 0 을 지우고 = 로 시작하는 수식을 넣습니다."),
        ("p", "③ 노란 칸이 여러 줄이면 — 첫 칸에만 수식을 넣고, 채우기 핸들(칸 오른쪽 아래 ■)을 마지막 노란 칸까지 끌어내립니다."),
        ("p", "④ 오른쪽 [채점] 칸이 ✔ 맞음 / ✘ 다시 로 스스로 바뀝니다. 수식이 달라도 결과가 같으면 맞음입니다."),
        ("p", "⑤ [채점표] 시트에서 전체 진행 상황을, 맨 끝 [정답] 시트에서 모범답안을 볼 수 있습니다."),
        ("note", "사람마다 결과가 달라야 정상입니다. 첫 칸은 맞았는데 아래 칸이 틀린다면, 끌어내리면서 범위나 기준 칸도 "
                 "같이 내려간 것입니다 — 움직이면 안 되는 주소에 F4 로 $ 를 붙이세요."),
        ("note", "값이 안 바뀌어 보이면 Ctrl + Alt + F9 (전체 다시 계산). 숫자를 손으로 적으면 맞음이 나오지 않습니다."),
    ])
    score = wb.create_sheet("채점표")
    made = []
    for no, p in enumerate(probs, 1):
        ws, mark, first, last = problem_sheet(wb, no, p)
        made.append((no, p, ws.title, mark, first, last))

    widths(score, {"A": 7, "B": 34, "C": 16, "D": 18})
    title(score, 1, "채점표 — %s %d문제" % (short, len(probs)), span=4)
    hdr(score, 3, ["번호", "문제", "넣을 칸", "채점"])
    for i, (no, p, sname, mark, first, last) in enumerate(made):
        r = 4 + i
        put(score, r, [no, p["title"], first if first == last else "%s → %s" % (first, last), None])
        score.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
        score.cell(row=r, column=3).alignment = Alignment(horizontal="center", vertical="center")
        c = score.cell(row=r, column=4, value="='%s'!%s" % (sname, mark))
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = Font(bold=True)
    end = 4 + len(made) - 1
    mark_rules(score, "D4:D%d" % end)
    tot = score.cell(row=end + 2, column=2, value="맞은 문제")
    tot.font = Font(bold=True)
    t = score.cell(row=end + 2, column=4,
                   value='=COUNTIF(D4:D%d,"*맞음*")&" / %d"' % (end, len(made)))
    t.font = Font(bold=True, size=13, color=NAVY)
    t.alignment = Alignment(horizontal="center")

    ans = wb.create_sheet("정답")
    widths(ans, {"A": 6, "B": 28, "C": 16, "D": 52, "E": 44})
    title(ans, 1, "정답 — 먼저 스스로 풀어 본 뒤에 보세요", span=5)
    note(ans, 2, "[넣을 칸] 의 첫 칸에 이 수식을 넣고, 칸이 여러 개면 끝까지 끌어내리면 됩니다.", span=5)
    hdr(ans, 4, ["번호", "문제", "넣을 칸", "모범답안 (첫 칸)", "힌트"])
    for i, (no, p, sname, mark, first, last) in enumerate(made):
        r = 5 + i
        put(ans, r, [no, p["title"], first if first == last else "%s:%s" % (first, last), None, p["hint"]],
            fill=GREEN if no % 2 == 0 else None)
        ans.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ans.cell(row=r, column=5).alignment = Alignment(vertical="center", wrap_text=True)
        lit(ans, r, 4, p["answer"], bold=True, color="1F4E79")
        ans.row_dimensions[r].height = 30
    if cat == "배열수식":
        note(wb["읽어보기"], 14, "배열 수식은 수식을 다 쓴 뒤 Enter 가 아니라 Ctrl + Shift + Enter 로 마칩니다. "
                                "수식 양옆에 { } 가 저절로 붙으면 제대로 들어간 것입니다(직접 치면 안 됩니다).", span=2)
    return save(wb, os.path.join(outdir, "fx" if GRADE == "2" else "fx1", dlname(idx, short)))


if __name__ == "__main__":
    outdir = sys.argv[1] if len(sys.argv) > 1 else "out"
    if len(sys.argv) > 2 and sys.argv[2] == "1":
        GRADE, CATS = "1", CATS1
    data = json.load(open(os.path.join(outdir, "web_problems.json" if GRADE == "2" else "web_problems_1.json"), encoding="utf-8"))
    for i, (cat, key, short) in enumerate(CATS, 1):
        probs = [p for p in data if p["cat"] == cat]
        build(cat, i, key, short, probs, outdir)
